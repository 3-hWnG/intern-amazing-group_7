"""Đọc tệp người dùng tải lên và đổi sang CẤU TRÚC ĐỊNH SẴN (NV3, 4A: không cần người dùng xác nhận).

Một bản ghi = {title, fields {tên trường: giá trị}, text (chữ để tìm), source (tệp · trang tính · dòng / trang)}.
NV5 (A5): tiêu đề cột hai tầng (ô gộp "Điểm" trên "Toán | Văn"), bảng nằm ngang (tên trường ở cột đầu), nhiều bảng trong
một trang tính (cách nhau bởi dòng trống), và người dùng tự chọn dòng tiêu đề trong "Cách đọc" (overrides).
- Bảng (CSV, Excel, JSON danh sách, bảng trong Word): code chọn cột tiêu đề, AI kiểm lại, code duyệt.
  Không khớp được -> phương án C: mỗi dòng thành một đoạn chữ "cột: giá trị; …" (kind = rows).
- Văn bản (TXT/MD, đoạn văn Word, PDF): chia theo MỤC (tiêu đề # / style Heading của Word / dòng trông như tiêu đề), rồi mỗi
  mục thành các đoạn <= CHUNK_CHARS ký tự cắt ở dòng trống > dòng > câu > khoảng trắng; mỗi đoạn mở đầu bằng tên mục (kind = text).
  PDF: nối các trang trước khi chia (mục vắt qua hai trang vẫn là một mục), nguồn ghi khoảng trang.
"""
from __future__ import annotations
import csv
import datetime
import io
import json
import re
from pathlib import Path

from . import llm, settings

READER_VERSION = 5   # tăng khi đổi cách đọc tệp: dữ liệu đọc bằng bản cũ được xử lý lại khi khởi động server
                     # 5 (19/10): văn bản chia theo MỤC (PDF nối trang + nhận dòng tiêu đề), không cắt giữa chữ
csv.field_size_limit(2**31 - 1)   # ô rất dài (mô tả, văn bản dán vào) không làm hỏng việc đọc CSV
SUPPORTED = {".csv", ".tsv", ".xlsx", ".xlsm", ".json", ".txt", ".md", ".docx", ".pdf"}
_TITLE_HINTS = ("tên", "ten", "name", "title", "tiêu đề", "tieu de", "sản phẩm", "san pham", "thủ tục", "thu tuc", "mặt hàng",
                "câu hỏi", "cau hoi", "question", "chủ đề", "chu de", "dịch vụ", "dich vu", "họ tên", "ho ten", "mục", "item")


class IngestError(Exception):
    pass


# ------------------------------------------------------------------ đọc tệp
def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()[:10] if isinstance(v, datetime.datetime) and not (v.hour or v.minute) else v.isoformat()
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return " ".join(str(v).split())


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1258", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read(path: Path) -> list[tuple]:
    """Trả danh sách phần: ("table", tên, rows[list[list[str]]]) hoặc ("text", tiêu đề, chữ, nguồn)."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise IngestError(f"Chưa hỗ trợ tệp {ext or 'không có đuôi'}. Hỗ trợ: CSV, Excel (.xlsx), JSON, TXT/MD, Word (.docx), PDF.")
    if ext in (".csv", ".tsv"):
        text = _decode(path.read_bytes())
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if ext == ".tsv" else csv.excel
        return [("table", path.stem, [[_cell(c) for c in r] for r in csv.reader(io.StringIO(text), dialect)])]
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        small = path.stat().st_size < 15 * 1024 * 1024   # tệp nhỏ: đọc đủ để biết ô gộp (tiêu đề hai tầng); tệp lớn: đọc nhanh
        wb = openpyxl.load_workbook(path, read_only=not small, data_only=True)
        out = []
        for ws in wb.worksheets:
            merged = [(m.min_row, m.min_col - 1, m.max_row, m.max_col - 1) for m in ws.merged_cells.ranges] if small else []
            out.append(("table", ws.title, [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)], merged))
        wb.close()
        return out
    if ext == ".json":
        try:
            data = json.loads(_decode(path.read_bytes()))
        except ValueError as e:
            raise IngestError(f"Tệp JSON hỏng: {e}") from None
        if isinstance(data, dict):
            lst = next((v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)), None)
            data = lst if lst is not None else data
        if isinstance(data, list) and data and all(isinstance(x, dict) for x in data):
            keys = list(dict.fromkeys(k for x in data for k in x))
            return [("table", path.stem, [keys] + [[_cell(x.get(k)) for k in keys] for x in data])]
        return [("text", path.stem, json.dumps(data, ensure_ascii=False, indent=1), path.name)]
    if ext in (".txt", ".md"):
        return _sections(_decode(path.read_bytes()), path.name, markdown=True)
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        out, buf = [], []
        body = d.element.body
        tables = {t._tbl: t for t in d.tables}
        paras = {p._p: p for p in d.paragraphs}
        is_head = lambda p: (p.style.name or "").lower().startswith(("heading", "title")) and p.text.strip()
        styled = any(is_head(p) for p in d.paragraphs)   # không có style Heading nào: đoán dòng tiêu đề như PDF
        for child in body.iterchildren():
            if child in paras:
                p = paras[child]
                if styled and is_head(p):
                    buf.append(HEAD_MARK + p.text.strip())
                elif p.text.strip():
                    buf.append(p.text.strip())
            elif child in tables:
                rows = [[_cell(c.text) for c in r.cells] for r in tables[child].rows]
                if len(rows) >= 2:
                    out += _split_sections(buf, path.stem, path.name, guess=not styled)
                    buf = []
                    out.append(("table", f"{path.stem} · bảng {sum(1 for x in out if x[0] == 'table') + 1}", rows))
        return out + _split_sections(buf, path.stem, path.name, guess=not styled)
    if ext == ".pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(str(path))
        except Exception as e:
            raise IngestError(f"Không đọc được PDF: {e}") from None
        lines, pages = [], []
        for i, page in enumerate(reader.pages, 1):
            for ln in (page.extract_text() or "").splitlines():
                lines.append(ln)
                pages.append(i)
        if not any(ln.strip() for ln in lines):
            raise IngestError("PDF không có chữ (có thể là ảnh quét). Hiện chỉ hỗ trợ PDF có chữ.")
        return _split_sections(lines, path.stem, path.name, guess=True, pages=pages)
    raise IngestError("Định dạng không hỗ trợ")


def _sections(text: str, source: str, markdown: bool) -> list[tuple]:
    """Chia văn bản theo tiêu đề Markdown (#) nếu có; không có # thì đoán dòng tiêu đề (looks_heading)."""
    lines, marked = [], False
    for line in text.splitlines():
        m = re.match(r"^\s{0,3}#{1,6}\s+(.+)$", line) if markdown else None
        marked = marked or bool(m)
        lines.append(HEAD_MARK + m.group(1).strip() if m else line)
    return _split_sections(lines, Path(source).stem, source, guess=not marked)


HEAD_MARK = "\x00H"   # dòng đã biết chắc là tiêu đề (Markdown #, style Heading của Word)
_LEGAL = re.compile(r"^(chương|mục|phần|điều|phụ lục)\s+([0-9]+[a-z]?|[ivxlc]+)(\s*$|[.:)])", re.I)   # "Điều 5 Luật…" (câu) thì không
_ROMAN = re.compile(r"^[IVXLC]{1,6}[.)]\s+\S")


def looks_heading(line: str) -> bool:
    """Dòng (của PDF / TXT / Word không có style) trông như tên mục: "Chương II", "Điều 5. Đối tượng…", "IV. HỒ SƠ",
    dòng IN HOA ngắn, hoặc dòng ngắn kết thúc bằng ":" ("Deployment:", "Implement Websearch:"). Câu thường, gạch đầu dòng,
    "1. …" (thường là ý liệt kê) thì không."""
    t = " ".join(line.split())
    if not t or len(t) > 90 or t[0] in "-*•+–(":
        return False
    if _LEGAL.match(t) or _ROMAN.match(t):
        return True
    letters = [c for c in t if c.isalpha()]
    if len(letters) >= 6 and len(t.split()) >= 2 and all(c.isupper() for c in letters) and not t.endswith((",", ";")):
        return True
    return t.endswith(":") and len(t) <= 60 and t.count(":") == 1 and len(t.split()) <= 8


def _split_sections(lines: list[str], stem: str, filename: str, guess: bool, pages: list[int] | None = None) -> list[tuple]:
    """Dòng -> các mục ("text", tên mục (không có thì tên tệp), chữ, nguồn, tên mục). Nhãn liền nhau (mục rỗng, vd. "Task:" rồi
    "Sửa code:") gộp thành một tên "Task: › Sửa code:". Dòng tiêu đề tự nó là một câu ("Điều 1. Giờ làm việc từ 8 giờ.") mà
    không có nội dung theo sau thì chính nó là nội dung, không bị bỏ."""
    out, head, last, buf, pg = [], "", "", [], []

    def label(h: str) -> bool:   # chỉ là nhãn (không mang nội dung): "Task:", "Chương II", "Điều 5. Đối tượng tuyển sinh"
        return not h.rstrip().endswith((".", "!", "?")) and len(h) <= 60

    def flush():
        body, h = "\n".join(buf).strip(), head
        if not body and h and not label(last):   # tiêu đề mang nội dung, không có thân
            body, h = last, h.rsplit(" › ", 1)[0] if " › " in h else ""
        if body:
            src = filename
            if pages and pg:
                a, b = min(pg), max(pg)
                src = f"{filename} · trang {a}" + (f"–{b}" if b != a else "")
            out.append(("text", h or stem, body, src, h))

    for i, line in enumerate(lines):
        known = line.startswith(HEAD_MARK)
        text = line[len(HEAD_MARK):] if known else line
        if known or (guess and looks_heading(text)):
            if "".join(buf).strip() or (head and not label(last)):
                flush()
                head = ""
            elif head:   # mục rỗng: nối tên với tiêu đề con
                head += " › "
            last = " ".join(text.split())
            head += last
            buf, pg = [], [pages[i]] if pages else []
        else:
            buf.append(text)
            if pages and text.strip():
                pg.append(pages[i])
    flush()
    return out


# --------------------------------------------------------- bảng -> bản ghi
def _numeric(v: str) -> bool:
    return bool(re.fullmatch(r"[\d\s.,%:/\-+đ₫$]*", v or ""))


def _header_index(rows: list[list[str]]) -> int | None:
    """Dòng tiêu đề cột = dòng ĐẦU TIÊN (trong 15 dòng đầu) mà các ô là nhãn chữ ngắn và phủ >= một nửa số cột.
    Dòng tên bảng / ghi chú phía trên (vd. "DANH SÁCH HỌC SINH LỚP ..." chỉ chiếm vài ô) bị bỏ qua.
    Gặp dòng toàn số trước khi thấy tiêu đề -> bảng không có tiêu đề cột."""
    width = max((sum(1 for c in r if c) for r in rows[:30]), default=0)
    fallback = None
    for i, r in enumerate(rows[:15]):
        filled = [c for c in r if c]
        if len(filled) < 2:
            continue
        labels = [c for c in filled if not _numeric(c) and len(c) <= 60]
        if len(labels) < max(2, len(filled) * 0.6):
            return fallback   # dòng dữ liệu (nhiều số) -> dừng
        if len(filled) >= width * 0.5:
            return i
        fallback = i if fallback is None else fallback
    return fallback


def _norm(h: str) -> str:
    return " ".join(_strip(h).lower().split())


def _strip(t: str) -> str:
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", t or "") if unicodedata.category(c) != "Mn").replace("đ", "d").replace("Đ", "D")


def _name_pair(headers: list[str]) -> tuple[int, int] | None:
    """Danh sách người kiểu Việt Nam hay tách "Họ và tên" thành 2 cột: họ + tên đệm | tên. Trả (cột họ, cột tên)."""
    for i in range(len(headers) - 1):
        if _norm(headers[i + 1]) == "ten" and _norm(headers[i]).startswith(("ho", "cot")):
            return i, i + 1
    return None


def _two_row_header(top: list[str], sub: list[str], spans: list[tuple[int, int]]) -> list[str] | None:
    """Tiêu đề hai tầng: dòng trên có nhãn nhóm phủ nhiều cột (ô gộp, vd. "Điểm"), dòng dưới có nhãn từng cột ("Toán", "Văn").
    spans = các khoảng cột gộp ở dòng trên (từ Excel). Không có thông tin ô gộp (CSV, tệp lớn): chỉ nhận khi nhãn nhóm
    phủ >= 3 cột, để cặp "Họ | (trống)" quen thuộc không bị nhầm. Trả tên cột đã ghép hoặc None."""
    filled = [c for c in sub if c]
    if len(filled) < 2 or any(_numeric(c) for c in filled) or any(len(c) > 60 for c in filled):
        return None   # dòng dưới là dữ liệu, không phải nhãn
    spans = [(a, b) for a, b in spans if b > a]
    if not spans:
        j = 0
        while j < len(top):
            if top[j]:
                k = j + 1
                while k < len(top) and not top[k] and sub[k]:
                    k += 1
                if k - j >= 3 and sub[j]:
                    spans.append((j, k - 1))
                j = k
            else:
                j += 1
    spans = [(a, b) for a, b in spans if any(sub[j] for j in range(a, min(b, len(sub) - 1) + 1))]
    if not spans:
        return None
    names = list(top)
    for a, b in spans:
        for j in range(a, min(b, len(names) - 1) + 1):
            if sub[j]:
                names[j] = f"{top[a]} - {sub[j]}" if top[a] else sub[j]
    for j in range(len(names)):
        if not names[j] and sub[j]:
            names[j] = sub[j]
    return names


def _clean_table(rows: list[list[str]], merged: list | None = None, header_row: int | None = None,
                 offset: int = 0) -> tuple[list[str], list[list[str]], dict]:
    """Trả (tên cột, dòng dữ liệu, cách đọc: dòng tiêu đề trong tệp, các dòng bỏ qua phía trên).
    header_row = dòng tiêu đề người dùng chọn (số dòng trong trang tính); offset = số dòng phía trên phần bảng này."""
    numbered = [(n, r) for n, r in enumerate(rows, offset + 1) if any(c for c in r)]   # số dòng thật trong tệp
    rows = [r for _, r in numbered]
    if not rows:
        return [], [], {}
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    hi = next((i for i, (n, _) in enumerate(numbered) if n == header_row), None) if header_row else None
    chosen = hi is not None
    if not chosen:
        hi = _header_index(rows)
    how = {"header_row": numbered[hi][0] if hi is not None else None,
           "skipped_above": [" | ".join(c for c in r if c)[:160] for r in rows[:hi or 0]]}
    if chosen:
        how["header_chosen"] = True
    two = None
    if hi is not None and hi + 1 < len(rows):
        spans = [(a, b) for (r1, a, r2, b) in (merged or []) if r1 == numbered[hi][0] and b > a]
        two = _two_row_header(rows[hi], rows[hi + 1], spans)
    if hi is None:
        headers, data = [f"Cột {i + 1}" for i in range(width)], rows
    elif two:
        headers, data = two, rows[hi + 2:]
        how["two_row_header"] = True
    else:
        headers, data = rows[hi], rows[hi + 1:]
        headers = [h or ("Họ và tên đệm" if i + 1 < width and _norm(headers[i + 1]) == "ten" else f"Cột {i + 1}")
                   for i, h in enumerate(headers)]   # cột không nhãn ngay trước cột "Tên" = phần họ và tên đệm
        seen = {}
        for i, h in enumerate(headers):   # tên cột trùng -> thêm số
            if h in seen:
                seen[h] += 1; headers[i] = f"{h} ({seen[h]})"
            else:
                seen[h] = 1
    keep = [i for i in range(width) if sum(1 for r in data if r[i]) > max(0, len(data) * 0.02)]
    return [headers[i] for i in keep], [[r[i] for i in keep] for r in data], how


def _title_score(header: str, values: list[str]) -> float:
    vals = [v for v in values if v]
    if not vals or len(vals) < len(values) * 0.6:
        return -1
    if sum(_numeric(v) for v in vals) > len(vals) * 0.5:
        return -1
    avg = sum(len(v) for v in vals) / len(vals)
    spaced = sum(" " in v for v in vals) / len(vals)
    if (spaced < 0.3 and avg > 15) or sum(v.startswith(("http://", "https://")) for v in vals) > len(vals) * 0.3:
        return -1   # mã, chuỗi băm, đường dẫn: duy nhất nhưng không phải tiêu đề
    uniq = len(set(vals)) / len(vals)
    s = uniq * 2 + (1 if 3 <= avg <= 150 else -1)
    if any(h in header.lower() for h in _TITLE_HINTS):
        s += 2
    return s


def _ai_title(headers: list[str], data: list[list[str]], guess: str | None) -> str | None:
    """AI kiểm lại lựa chọn cột tiêu đề (chế độ JSON, ~1 giây). Lỗi/không hợp lệ -> None."""
    sample = "\n".join(" | ".join(r) for r in data[:4])
    from core import gpu
    try:
        gpu.scheduler.acquire(gpu.P3, timeout=120)   # nạp tệp: ưu tiên thấp nhất (G3.1)
    except gpu.Busy:
        return None
    try:
        out = _ai_title_call(headers, sample, guess)
    finally:
        gpu.scheduler.release()
    t = str(out.get("title_column", "")).strip()
    return t if t in headers else None


def _ai_title_call(headers: list[str], sample: str, guess: str | None) -> dict:
    return llm.chat_json([
        {"role": "system", "content": "Bạn chọn cột TIÊU ĐỀ cho một bảng dữ liệu: cột mô tả mỗi dòng là cái gì (tên sản phẩm, tên thủ tục, câu hỏi...). "
                                      "Chỉ chọn một tên cột có trong danh sách. Không có cột nào phù hợp thì để chuỗi rỗng."},
        {"role": "user", "content": f"Các cột: {headers}\nGợi ý của code: {guess or '(không có)'}\nVài dòng mẫu:\n{sample}"}],
        {"type": "object", "properties": {"title_column": {"type": "string"}}, "required": ["title_column"]}, timeout=30)


def _consistency(lines: list[list[str]]) -> list[float]:
    """Mỗi dòng (hoặc cột): tỉ lệ ô cùng kiểu (toàn số / toàn chữ). 1.0 = đồng nhất."""
    out = []
    for l in lines:
        v = [c for c in l if c]
        if len(v) >= 2:
            k = sum(_numeric(c) for c in v) / len(v)
            out.append(max(k, 1 - k))
    return out


def _sideways(headers: list[str], data: list[list[str]]) -> bool:
    """Bảng nằm ngang: cột đầu là tên trường ("Giá", "Màu"…), mỗi cột sau là một bản ghi. Nhận ra khi từng DÒNG cùng kiểu
    (dòng "Giá" toàn số) còn từng CỘT lẫn kiểu — ngược với bảng thường. Chỉ xét bảng nhỏ có cả dòng số và dòng chữ."""
    if not 2 <= len(data) <= 40 or len(headers) < 3:
        return False
    labels = [r[0] for r in data]
    if not all(l and not _numeric(l) and len(l) <= 60 for l in labels) or len(set(labels)) != len(labels):
        return False
    body = [r[1:] for r in data]
    rows_c, cols_c = _consistency(body), _consistency([list(c) for c in zip(*body)])
    if not rows_c or not cols_c:
        return False
    kinds = [sum(_numeric(c) for c in l if c) / max(1, sum(1 for c in l if c)) for l in body]
    return (sum(rows_c) / len(rows_c) >= 0.9 and sum(cols_c) / len(cols_c) <= 0.75
            and any(k >= 0.8 for k in kinds) and any(k <= 0.2 for k in kinds))


def table_records(name: str, rows: list[list[str]], filename: str, use_ai: bool = True, merged: list | None = None,
                  header_row: int | None = None, offset: int = 0) -> tuple[list[dict], dict]:
    headers, data, how = _clean_table(rows, merged, header_row, offset)
    if not data:
        return [], {"kind": "empty"}
    if _sideways(headers, data):   # xoay lại: mỗi cột thành một bản ghi
        corner = headers[0] if headers[0] and not headers[0].startswith("Cột") else "Tên"
        headers, data = [corner] + [r[0] for r in data], [[h] + [r[i] for r in data] for i, h in enumerate(headers) if i > 0]
        how["sideways"] = True
    if len(headers) == 1:   # một cột: là ghi chú/văn bản (câu dài) thì coi như văn bản
        vals = [headers[0]] + [r[0] for r in data]
        if sum(len(v) for v in vals) / len(vals) > 40 or headers[0].startswith("Cột"):
            recs = text_records(name, "\n\n".join(vals), f"{filename} · {name}")
            return recs, {"kind": "text", "title": name, "chunks": len(recs)}
    scores = {h: _title_score(h, [r[i] for r in data[:500]]) for i, h in enumerate(headers)}
    guess = max(scores, key=scores.get) if scores and max(scores.values()) > 0 else None
    title = (_ai_title(headers, data, guess) if use_ai else None) or guess
    if title and scores.get(title, -1) < 0:   # AI chọn cột không dùng được (trống / toàn số): giữ lựa chọn của code
        title = guess
    recs = []
    pair = _name_pair(headers)
    if title and pair and headers.index(title) in pair:   # tiêu đề = họ tên đầy đủ (ghép 2 cột)
        title_label = f"{headers[pair[0]]} + {headers[pair[1]]}"
    else:
        pair, title_label = None, title
    if title:
        ti = headers.index(title)
        for n, r in enumerate(data, 1):
            fields = {h: v for h, v in zip(headers, r) if v and h != title}
            t = (" ".join(x for x in (r[pair[0]], r[pair[1]]) if x) if pair else r[ti]) or f"Dòng {n}"
            if pair:
                fields = {h: v for h, v in zip(headers, r) if v}
            recs.append({"title": t[:300], "fields": fields,
                         "text": f"{'Họ và tên' if pair else title}: {t}\n" + "\n".join(f"{k}: {v}" for k, v in fields.items()),
                         "source": f"{filename} · {name} · dòng {n}"})
        return recs, {"kind": "table", "title_column": title_label, "fields": [h for h in headers if h != title], "sheet": name, **how}
    for n, r in enumerate(data, 1):   # phương án C: không có cột tiêu đề dùng được
        fields = {h: v for h, v in zip(headers, r) if v}
        recs.append({"title": f"{name} · dòng {n}", "fields": fields,
                     "text": "; ".join(f"{k}: {v}" for k, v in fields.items()), "source": f"{filename} · {name} · dòng {n}"})
    return recs, {"kind": "rows", "sheet": name, "fields": headers, **how}


def _cut(s: str, size: int) -> list[str]:
    """Một dòng dài hơn size: cắt ở cuối câu (". ", "; ", ": "), rồi ", ", cuối cùng ở khoảng trắng — không cắt giữa chữ."""
    out = []
    while len(s) > size:
        cut = -1
        for sep in (". ", "? ", "! ", "; ", ": ", ", ", " "):
            k = s.rfind(sep, 0, size)
            if k > size * (0.5 if sep != " " else 0.2):
                cut = k + len(sep)
                break
        cut = cut if cut > 0 else size   # một "chữ" dài hơn size (đường link…): đành cắt cứng
        out.append(s[:cut].strip())
        s = s[cut:].strip()
    if s:
        out.append(s)
    return out


def text_records(title: str, text: str, source: str, head: str = "") -> list[dict]:
    """Chia một mục văn bản thành đoạn <= CHUNK_CHARS ký tự: gom theo đoạn văn (dòng trống / gạch đầu dòng); đoạn văn dài thì
    gom theo DÒNG (chữ PDF không có dòng trống); dòng dài thì cắt ở câu / khoảng trắng. Gối đầu đơn vị cuối nếu nó ngắn.
    head (tên mục) đứng đầu mỗi đoạn để tìm kiếm và AI biết đoạn thuộc mục nào."""
    size = settings.get("CHUNK_CHARS")
    lead = f"{head}\n" if head else ""
    room = max(200, size - len(lead))
    units = []
    for p in re.split(r"\n\s*\n|\n(?=[-*•]\s)", text):
        p = p.strip()
        if not p:
            continue
        if len(p) <= room:
            units.append(p)
            continue
        for ln in p.split("\n"):
            units += _cut(ln.strip(), room) if ln.strip() else []
    pieces, buf = [], []
    for u in units:
        if buf and sum(len(x) + 1 for x in buf) + len(u) > room:
            pieces.append("\n".join(buf))
            buf = [buf[-1]] if len(buf[-1]) < room * 0.25 and len(buf[-1]) + 1 + len(u) <= room else []
        buf.append(u)
    if buf:
        pieces.append("\n".join(buf))
    multi = len(pieces) > 1
    return [{"title": f"{title} (phần {i})" if multi else title, "fields": {}, "text": lead + t, "source": source}
            for i, t in enumerate(pieces, 1)]


def _blocks(rows: list[list[str]]) -> list[tuple[int, list[list[str]]]]:
    """Nhiều bảng trong một trang tính: tách tại dòng trống khi phần sau bắt đầu bằng dòng tiêu đề riêng.
    Phần quá nhỏ (dòng tên bảng phía trên) gộp vào phần sau; dòng trống giữa dữ liệu thì gộp vào phần trước.
    Trả [(số dòng phía trên phần này, các dòng)]."""
    segs, cur, start = [], [], 0
    for i, r in enumerate(rows):
        if any(c for c in r):
            if not cur:
                start = i
            cur.append(r)
        elif cur:
            segs.append((start, cur)); cur = []
    if cur:
        segs.append((start, cur))
    out: list[tuple[int, list[list[str]]]] = []
    pending = None
    for start, seg in segs:
        if pending is not None:   # phần nhỏ phía trên (tên bảng) gộp vào phần này
            seg = rows[pending:start] + seg
            start, pending = pending, None
        if len(seg) < 3:
            pending = start
            continue
        width = max(sum(1 for c in r if c) for r in seg[:5])
        is_new = out and _header_index(seg) == 0 and width >= 2 and [c for c in seg[0] if c] != [c for c in out[-1][1][0] if c]
        if out and not is_new:
            p0, prev = out[-1]
            out[-1] = (p0, rows[p0:start] + seg)
        else:
            out.append((start, seg))
    if pending is not None:
        if out:
            p0, prev = out[-1]
            out[-1] = (p0, rows[p0:])
        else:
            out.append((pending, rows[pending:]))
    return out or [(0, rows)]


def to_records(path: Path, filename: str, use_ai: bool = True, overrides: dict | None = None) -> tuple[list[dict], dict]:
    """Tệp -> (bản ghi theo cấu trúc định sẵn, mô tả cách đã khớp). overrides = {tên phần: dòng tiêu đề người dùng chọn}."""
    recs, maps = [], []
    overrides = overrides or {}
    for part in read(path):
        if part[0] == "table":
            merged = part[3] if len(part) > 3 else None
            blocks = _blocks(part[2])
            for k, (off, rows) in enumerate(blocks, 1):
                pname = part[1] if len(blocks) == 1 else f"{part[1]} · bảng {k}"
                r, m = table_records(pname, rows, filename, use_ai, merged, overrides.get(pname), off)
                if r:
                    recs += r; maps.append(m if m["kind"] != "text" else {"kind": "text", "title": m["title"], "chunks": m["chunks"]})
        else:
            r = text_records(part[1], part[2], part[3], part[4] if len(part) > 4 else "")
            if r:
                recs += r; maps.append({"kind": "text", "title": part[1], "chunks": len(r)})
    if not recs:
        raise IngestError("Không tìm thấy nội dung nào trong tệp.")
    kinds = {m["kind"] for m in maps}
    kind = "table" if kinds == {"table"} else ("text" if kinds == {"text"} else ("rows" if kinds == {"rows"} else "mixed"))
    return recs, {"kind": kind, "parts": maps, "reader": READER_VERSION}


def describe(mapping: dict, n: int) -> str:
    """Câu mô tả cho người dùng: đã khớp cấu trúc thế nào."""
    parts = mapping.get("parts", [])
    tables = [m for m in parts if m["kind"] == "table"]
    rows = [m for m in parts if m["kind"] == "rows"]
    texts = [m for m in parts if m["kind"] == "text"]
    out = []
    if tables:
        out.append("Bảng: cột tiêu đề " + ", ".join(f"\"{m['title_column']}\"" for m in tables[:3])
                   + f" ({sum(len(m['fields']) for m in tables)} trường)")
    if rows:
        out.append(f"{len(rows)} bảng không khớp được cột tiêu đề, lưu dạng dòng chữ")
    if texts:
        out.append(f"Văn bản: {sum(m['chunks'] for m in texts)} đoạn")
    return "; ".join(out) + f" · {n} bản ghi"

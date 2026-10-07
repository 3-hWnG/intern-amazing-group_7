"""NV3: tải dữ liệu -> cấu trúc định sẵn -> chỉ mục -> Chuyên gia trả lời từ dữ liệu, có nguồn.
Model giả (embedding, reranker, LLM) để chạy nhanh, không cần GPU/Ollama.
Chạy (từ thư mục server/): set PYTHONPATH=<pyroot> && python ../system4/tests/test_nv3.py"""
import hashlib
import io
import json
import math
import os
import re
import sys
import tempfile
import time

SERVER = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "server")
sys.path.insert(0, SERVER)
TMP = tempfile.mkdtemp()
os.environ.update(S3_DB_PATH=os.path.join(TMP, "s3.db"), S3_USE_LLM="0", S4_ENABLED="1", S4_WARMUP="0", S4_RUNTIME_DIR=os.path.join(TMP, "s4"))

from fastapi.testclient import TestClient
import main
from system3.system4.server import auth, datasets, db, ingest, llm, search


def fake_embed(texts):
    """Túi từ băm vào 1024 chiều: câu chung nhiều từ thì gần nhau (đủ để kiểm đường đi, không phải chất lượng)."""
    out = []
    for t in texts:
        v = [0.0] * 1024
        for w in re.findall(r"\w+", t.lower()):
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % 1024] += 1
        n = math.sqrt(sum(x * x for x in v)) or 1
        out.append([x / n for x in v])
    return out


def fake_rerank(query, texts):
    q = set(re.findall(r"\w+", query.lower()))
    return [len(q & set(re.findall(r"\w+", t.lower()))) * 3 - 4.0 for t in texts]   # trùng >= 1 từ quan trọng mới qua ngưỡng -3


prompts, fast_answers = [], []


async def fake_fast(messages, schema):
    prompts.append(messages)
    d = fast_answers.pop(0) if fast_answers else {"plan": "", "answer": "Theo dữ liệu, giá là 120.000đ [1].", "ask_back": False,
                                                   "choices": [], "sources": [1]}
    if "sources" not in schema["properties"]:
        d = {k: v for k, v in d.items() if k != "sources"}
    raw = json.dumps(d)
    for i in range(0, len(raw), 7):
        yield raw[i:i + 7]


search.embed = fake_embed
search.rerank = fake_rerank
llm.stream_json = fake_fast
llm.chat_json = lambda *a, **k: {}


def csv_bytes(rows):
    s = io.StringIO()
    for r in rows:
        s.write(",".join(r) + "\n")
    return s.getvalue().encode("utf-8")


PRODUCTS = [["Tên sản phẩm", "Giá", "Mô tả", "Mã"]] + [
    ["Bánh mì thịt", "25.000đ", "Bánh mì kẹp thịt nướng, rau thơm", "SP001"],
    ["Phở bò tái", "55.000đ", "Phở bò Hà Nội, nước dùng hầm xương", "SP002"],
    ["Cà phê sữa đá", "29.000đ", "Cà phê phin pha sữa đặc", "SP003"],
    ["Trà đào cam sả", "45.000đ", "Trà đào tươi, cam, sả", "SP004"],
    ["Bún chả Hà Nội", "60.000đ", "Bún chả than hoa, nem rán", "SP005"],
]


def wait_ready(c, ds_id, timeout=30):
    t = time.time()
    while time.time() - t < timeout:
        d = next(x for x in c.get("/s4/datasets").json()["datasets"] if x["id"] == ds_id)
        if d["status"] in ("ready", "error"):
            return d
        time.sleep(0.2)
    raise AssertionError(f"dataset {ds_id} chưa xong")


with TestClient(main.app) as c:
    tok = {}
    for name in ("dev1", "lan", "minh"):
        c.cookies.clear()
        tok[name] = c.post("/s4/auth/signup", json={"username": name, "password": "123456"}).cookies.get(auth.COOKIE)

    def as_user(n):
        c.cookies.clear(); c.cookies.set(auth.COOKIE, tok[n])

    def chat(text, **kw):
        r = c.post("/s4/chat", json={"text": text, **kw})
        assert r.status_code == 200, r.text
        return [json.loads(l[6:]) for l in r.text.split("\n\n") if l.startswith("data: ")]

    # ---- chưa có dữ liệu: AI chung
    as_user("lan")
    info = c.get("/s4/datasets").json()
    assert info["datasets"] == [] and info["active"]["datasets"] == 0 and info["quota"]["limit"] == 1024 * 1024 * 1024
    ev = chat("Giá phở bò bao nhiêu?")
    assert ev[0]["specialist"] is False and "CHẾ ĐỘ CHUYÊN GIA" not in prompts[-1][0]["content"]

    # ---- tải lên: kiểm định dạng, cỡ tệp
    assert c.post("/s4/datasets", files={"file": ("virus.exe", b"MZ...", "application/octet-stream")}).status_code == 422
    assert c.post("/s4/datasets", files={"file": ("rong.csv", b"", "text/csv")}).status_code == 422
    r = c.post("/s4/datasets", files={"file": ("menu quán.csv", csv_bytes(PRODUCTS), "text/csv")})
    assert r.status_code == 200, r.text
    d = wait_ready(c, r.json()["dataset"]["id"])
    assert d["status"] == "ready" and d["kind"] == "table" and d["n_records"] == 5 and d["active"], d
    assert d["mapping"]["parts"][0]["title_column"] == "Tên sản phẩm", d["mapping"]
    assert "Bảng: cột tiêu đề \"Tên sản phẩm\"" in d["message"]
    menu_id = d["id"]

    # ---- Chuyên gia: tìm -> AI chỉ dùng dữ liệu -> nguồn
    ev = chat("Phở bò tái giá bao nhiêu?")
    assert ev[0]["specialist"] is True
    sp = prompts[-1][0]["content"]
    assert "CHẾ ĐỘ CHUYÊN GIA" in sp and "KHÔNG dùng kiến thức chung" in sp and "[1] Phở bò tái" in sp, sp
    cid = ev[0]["conversation_id"]
    m = c.get(f"/s4/conversations/{cid}/messages").json()["messages"][-1]
    assert m["meta"]["specialist"] and m["meta"]["sources"][0]["title"] == "Phở bò tái", m["meta"]
    assert m["meta"]["retrieval"]["reranked"] is True
    rec_id = m["meta"]["sources"][0]["record_id"]
    rec = c.get(f"/s4/records/{rec_id}").json()
    assert rec["record"]["fields"]["Giá"] == "55.000đ" and rec["dataset"]["name"] == "menu quán"
    # câu ngắn nối tiếp: tìm kèm câu hỏi trước
    chat("còn mô tả?", conversation_id=cid)
    assert "[1] Phở bò tái" in prompts[-1][0]["content"], "câu ngắn ghép ngữ cảnh câu trước để tìm"
    # không có trong dữ liệu -> báo không có, hỏi lại
    fast_answers.append({"plan": "", "answer": "Dữ liệu không có thông tin này. Bạn muốn hỏi món nào?", "ask_back": True,
                         "choices": ["Phở bò tái", "Bún chả Hà Nội"], "sources": []})
    ev = chat("Thời tiết Đà Lạt hôm nay thế nào?")
    assert "KHÔNG tìm thấy đoạn nào liên quan" in prompts[-1][0]["content"]
    m = c.get(f"/s4/conversations/{ev[0]['conversation_id']}/messages").json()["messages"][-1]
    assert m["meta"]["choices"] == ["Phở bò tái", "Bún chả Hà Nội"] and not m["meta"].get("sources")
    # tìm không thấy mà AI vẫn tự trả lời bằng kiến thức chung -> code thay bằng câu cố định (yêu cầu .docx)
    fast_answers.append({"plan": "", "answer": "Truyện Kiều là tác phẩm của Nguyễn Du, viết vào đầu thế kỷ 19, gồm 3.254 câu thơ lục bát kể về cuộc đời Thúy Kiều.",
                         "ask_back": False, "choices": [], "sources": []})
    ev = chat("Ai là tác giả truyện Kiều?")
    assert any(e["type"] == "replace" for e in ev), [e["type"] for e in ev]
    m = c.get(f"/s4/conversations/{ev[0]['conversation_id']}/messages").json()["messages"][-1]
    assert m["content"].startswith("Thông tin này không có trong các bộ dữ liệu") and m["meta"]["guard"] == "Chỉ trả lời từ dữ liệu", m
    fast_answers.append({"plan": "", "answer": "Dạ, không có gì ạ!", "ask_back": False, "choices": [], "sources": []})
    ev = chat("cảm ơn nhé")
    m = c.get(f"/s4/conversations/{ev[0]['conversation_id']}/messages").json()["messages"][-1]
    assert m["content"] == "Dạ, không có gì ạ!", "câu ngắn xã giao không bị thay"
    # không ghi [n] -> hiện "đã tra cứu"
    fast_answers.append({"plan": "", "answer": "Cà phê sữa đá có giá 29.000đ.", "ask_back": False, "choices": [], "sources": []})
    ev = chat("Cà phê sữa đá bao nhiêu tiền")
    m = c.get(f"/s4/conversations/{ev[0]['conversation_id']}/messages").json()["messages"][-1]
    assert m["meta"]["consulted"][0]["title"] == "Cà phê sữa đá", m["meta"]

    # ---- tắt bộ dữ liệu -> AI chung
    assert c.patch(f"/s4/datasets/{menu_id}", json={"active": False}).json()["active"]["datasets"] == 0
    assert chat("Phở bò tái giá bao nhiêu?")[0]["specialist"] is False
    assert c.patch(f"/s4/datasets/{menu_id}", json={"active": True}).status_code == 200

    # ---- tắt reranker: vẫn tìm được (từ khoá + nghĩa)
    as_user("dev1")
    c.post("/s4/settings", json={"values": {"RERANKER_ENABLED": False}})
    as_user("lan")
    ev = chat("Bún chả Hà Nội giá bao nhiêu")
    m = c.get(f"/s4/conversations/{ev[0]['conversation_id']}/messages").json()["messages"][-1]
    assert m["meta"]["retrieval"]["reranked"] is False and "Bún chả Hà Nội" in prompts[-1][0]["content"]
    as_user("dev1")
    c.post("/s4/settings", json={"values": {"RERANKER_ENABLED": True}})

    # ---- giới hạn: cảnh báo + giới hạn cứng
    c.post("/s4/settings", json={"values": {"ACTIVE_WARN_DATASETS": 1, "ACTIVE_MAX_DATASETS": 2, "MAX_UPLOAD_MB": 1}})
    as_user("lan")
    d2 = wait_ready(c, c.post("/s4/datasets", files={"file": ("ghi chu.txt", "Giờ mở cửa: 7h đến 22h hằng ngày.\n\nĐịa chỉ: 12 Lê Lợi.".encode(), "text/plain")}).json()["dataset"]["id"])
    assert d2["kind"] == "text" and d2["active"]
    s = c.get("/s4/datasets").json()["active"]
    assert s["datasets"] == 2 and s["warnings"], s
    d3 = wait_ready(c, c.post("/s4/datasets", files={"file": ("them.csv", csv_bytes(PRODUCTS[:3]), "text/csv")}).json()["dataset"]["id"])
    assert not d3["active"] and "vượt giới hạn" in d3["message"], d3
    assert c.patch(f"/s4/datasets/{d3['id']}", json={"active": True}).status_code == 422
    big = b"a,b\n" + b"x" * (1024 * 1024 + 10)
    assert c.post("/s4/datasets", files={"file": ("to.csv", big, "text/csv")}).status_code == 422, "user: tệp > MAX_UPLOAD_MB"
    as_user("dev1")
    r = c.post("/s4/datasets", files={"file": ("to.csv", big, "text/csv")})
    assert r.status_code == 200, "dev không giới hạn"
    assert c.get("/s4/datasets").json()["quota"]["limit"] is None

    # ---- quyền xem: chủ + dev (2B); người khác không
    as_user("minh")
    assert c.get(f"/s4/records/{rec_id}").status_code == 404
    assert c.get(f"/s4/datasets/{menu_id}/records").status_code == 404
    assert c.patch(f"/s4/datasets/{menu_id}", json={"active": False}).status_code == 404
    assert c.get("/s4/datasets").json()["datasets"] == []
    as_user("dev1")
    assert c.get(f"/s4/datasets/{menu_id}/records?q=phở").json()["total"] == 1
    as_user("lan")
    assert c.get(f"/s4/datasets/{menu_id}/records").json()["total"] == 5

    # ---- xoá: bản ghi, chỉ mục từ khoá, vector đều mất
    assert c.delete(f"/s4/datasets/{menu_id}").json()["ok"]
    assert db.run("SELECT COUNT(*) n FROM records WHERE dataset_id=?", (menu_id,), one=True)["n"] == 0
    assert db.run("SELECT COUNT(*) n FROM records_fts WHERE rowid NOT IN (SELECT id FROM records)", one=True)["n"] == 0
    found, _ = search.search(db.user_by_name("lan")["id"], "Phở bò tái giá bao nhiêu")
    assert all(x["dataset_id"] != menu_id for x in found)

    # ---- nạp lại việc dở khi khởi động
    db.update_dataset(d2["id"], status="processing", progress=40)
    datasets.resume_unfinished()
    assert wait_ready(c, d2["id"])["status"] == "ready"

# ---- đọc các định dạng (không qua web)
import docx
p = os.path.join(TMP, "quy_dinh.docx")
doc = docx.Document()
doc.add_heading("Chính sách đổi trả", level=1)
doc.add_paragraph("Khách được đổi trả trong 7 ngày nếu sản phẩm còn nguyên tem.")
t = doc.add_table(rows=3, cols=2)
for i, (a, b) in enumerate([("Hạng thành viên", "Ưu đãi"), ("Bạc", "Giảm 5%"), ("Vàng", "Giảm 10%")]):
    t.cell(i, 0).text, t.cell(i, 1).text = a, b
doc.save(p)
from pathlib import Path
recs, mp = ingest.to_records(Path(p), "quy_dinh.docx", use_ai=False)
assert any(r["title"] == "Chính sách đổi trả" for r in recs) and any(r["title"] == "Vàng" for r in recs), [r["title"] for r in recs]
assert mp["kind"] == "mixed"
pj = os.path.join(TMP, "x.json")
open(pj, "w", encoding="utf-8").write(json.dumps({"items": [{"name": "Gói Cơ bản", "price": 100}, {"name": "Gói Nâng cao", "price": 200}]}))
recs, mp = ingest.to_records(Path(pj), "x.json", use_ai=False)
assert [r["title"] for r in recs] == ["Gói Cơ bản", "Gói Nâng cao"] and recs[0]["fields"]["price"] == "100"
recs, mp = ingest.to_records(Path(os.path.join(TMP, "hash.csv")) if False else Path(pj), "x.json", use_ai=False)
ph = os.path.join(TMP, "hash.csv")
open(ph, "w", encoding="utf-8").write("ma,ten hang\n" + "\n".join(f"{hashlib.md5(str(i).encode()).hexdigest()},Hàng số {i} loại tốt" for i in range(20)))
recs, mp = ingest.to_records(Path(ph), "hash.csv", use_ai=False)
assert mp["parts"][0]["title_column"] == "ten hang", "cột mã băm không được làm tiêu đề"
try:
    ingest.to_records(Path(p).with_suffix(".xls"), "a.xls")
    raise AssertionError("phải báo không hỗ trợ")
except ingest.IngestError as e:
    assert "Chưa hỗ trợ" in str(e)

print("OK test_nv3")

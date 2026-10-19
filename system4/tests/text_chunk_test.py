"""Chia văn bản theo MỤC (READER_VERSION 5, 19/10/2026): PDF nối trang + nhận dòng tiêu đề, không cắt giữa chữ, mỗi đoạn
mở đầu bằng tên mục, AI nhận trọn đoạn. Dòng PDF mô phỏng theo đúng Kế hoạch tuần tới.pdf (chữ pypdf trả về: không có dòng trống).
Không cần model. Chạy:  python run_server.py system4/tests/text_chunk_test.py"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "server"))
TMP = tempfile.mkdtemp()
os.environ.update(S4_RUNTIME_DIR=os.path.join(TMP, "s4"))

from pathlib import Path
from system3.system4.server import ingest, settings

SIZE = settings.get("CHUNK_CHARS")

# --- nhận dòng tiêu đề
for t in ("Task:", "Implement Websearch:", "Deployment:", "Điều 5. Đối tượng tuyển sinh", "Chương II", "CHƯƠNG II. QUY ĐỊNH CHUNG",
          "IV. HỒ SƠ ĐĂNG KÝ", "Phụ lục 1", "HƯỚNG DẪN SỬ DỤNG"):
    assert ingest.looks_heading(t), t
for t in ("Hồng: Đừng Implement", "Điều 5 Luật Giáo dục quy định rằng", "- ý một:", "Deployment: Không làm hỗ trợ Anh An (Nhân)",
          "Lưu 1 -2 thủ tục thôi, 8 là quá nhiều", "Domain cứng: Không được xóa, luật xử lý riêng. Chỉ có dev add được.", "OK", ""):
    assert not ingest.looks_heading(t), t

# --- PDF: hai trang, mục "Implement Websearch:" vắt qua trang 2; dòng không có dòng trống (như pypdf)
long_line = "Chỉ con AI chỗ hệ thống websearch (System 1) trên máy mấy ông rồi nói nó làm 1 nút tương tự trên web mình đi " * 14
lines = ["Tổng hợp các yêu cầu thầy nêu", "Task:", "Sửa code:", "Lưu 1 -2 thủ tục thôi, 8 là quá nhiều",
         "Domain mềm: Friendly mode. Cả user và dev add được. Không có persistent dataset cho users",
         "Deployment:", "Đẩy ra docker, tách ra từng cục", "Implement Websearch:", long_line.strip(),
         "Websearch áp dụng cho tất cả domain.", "Ghi thêm: Không giải quyết vấn đề nạp dữ liệu"]
pages = [1] * 9 + [2, 2]
parts = ingest._split_sections(lines, "Kế hoạch", "Kế hoạch.pdf", guess=True, pages=pages)
heads = [p[4] for p in parts]
assert heads == ["", "Task: › Sửa code:", "Deployment:", "Implement Websearch:"], heads   # tiêu đề liền nhau gộp tên
assert parts[0][1] == "Kế hoạch" and parts[1][1] == "Task: › Sửa code:" and parts[1][3] == "Kế hoạch.pdf · trang 1"
assert parts[3][3] == "Kế hoạch.pdf · trang 1–2", parts[3][3]                               # mục vắt qua hai trang: một mục
recs = []
for p in parts:
    recs += ingest.text_records(p[1], p[2], p[3], p[4])
dom = next(r for r in recs if "Domain mềm" in r["text"])
assert dom["text"].startswith("Task: › Sửa code:\n"), dom["text"][:40]                      # mỗi đoạn mang tên mục
web = [r for r in recs if r["title"].startswith("Implement Websearch:")]
assert len(web) >= 2 and all(r["text"].startswith("Implement Websearch:\n") for r in web), [r["title"] for r in web]
words = set(long_line.split())
for r in recs:
    assert len(r["text"]) <= SIZE, (r["title"], len(r["text"]))                              # không đoạn nào dài quá CHUNK_CHARS
    body = r["text"].split("\n", 1)[-1]
    for w in (body.split()[0], body.split()[-1]):                                           # không cắt giữa chữ
        assert w in words or w in " ".join(lines), (r["title"], w)
assert any("Websearch áp dụng cho tất cả domain." in r["text"] for r in web)

# --- dòng tiêu đề tự nó là nội dung ("Điều 1. Giờ làm việc…" rồi ngay "Điều 2."): không mất chữ, mỗi điều một mục
law = ["Quy chế nội bộ", "Chương I", "Điều 1. Giờ làm việc từ 8 giờ đến 17 giờ.", "Điều 2. Nghỉ trưa một tiếng.",
       "Điều 3. Đi muộn", "Quá ba lần trong tháng thì bị nhắc nhở."]
parts = ingest._split_sections(law, "qc", "qc.txt", guess=True)
got = [(p[1], p[2]) for p in parts]
assert got == [("qc", "Quy chế nội bộ"), ("Chương I", "Điều 1. Giờ làm việc từ 8 giờ đến 17 giờ."), ("qc", "Điều 2. Nghỉ trưa một tiếng."),
               ("Điều 3. Đi muộn", "Quá ba lần trong tháng thì bị nhắc nhở.")], got

# --- đoạn dài một dòng không có dấu câu: cắt ở khoảng trắng
pieces = ingest._cut("chữ " * 400, 300)
assert all(len(x) <= 300 and not x.endswith("ch") for x in pieces) and " ".join(pieces).split() == ["chữ"] * 400

# --- TXT: có # thì theo #, không có # thì đoán dòng tiêu đề
md = Path(TMP) / "a.md"
md.write_text("# Giới thiệu\nDòng một.\nTask:\nkhông phải tiêu đề vì đã có #\n# Phí\nPhí là 20.000đ.\n", encoding="utf-8")
assert [p[4] for p in ingest.read(md)] == ["Giới thiệu", "Phí"]
txt = Path(TMP) / "b.txt"
txt.write_text("ĐIỀU KIỆN THAM GIA\nCó doanh nghiệp.\nHồ sơ:\nĐơn đăng ký.\n", encoding="utf-8")
assert [p[4] for p in ingest.read(txt)] == ["ĐIỀU KIỆN THAM GIA", "Hồ sơ:"]

# --- Word không có style Heading: đoán dòng tiêu đề; có style thì chỉ theo style
import docx
d = docx.Document()
for t in ("Quy định chung:", "Áp dụng cho mọi người.", "Lệ phí:", "Miễn phí."):
    d.add_paragraph(t)
p1 = Path(TMP) / "c.docx"
d.save(str(p1))
assert [p[4] for p in ingest.read(p1)] == ["Quy định chung:", "Lệ phí:"]
d = docx.Document()
d.add_heading("Mục A", 1)
d.add_paragraph("Lệ phí:")
d.add_paragraph("Miễn phí.")
p2 = Path(TMP) / "d.docx"
d.save(str(p2))
assert [p[4] for p in ingest.read(p2)] == ["Mục A"]

# --- AI nhận trọn đoạn: không cắt ngắn hơn CHUNK_CHARS
assert settings.evidence_chars() >= SIZE
assert ingest.READER_VERSION >= 5
print("OK text_chunk_test")

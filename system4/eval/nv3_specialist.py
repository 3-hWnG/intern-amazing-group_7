"""NV3 — đo chế độ Chuyên gia trên model THẬT (qua web đang chạy, S4_ENABLED=1). Chấm bằng luật.

    python system4/eval/nv3_specialist.py http://127.0.0.1:8399

Tải lên MAU_100_THU_TUC.xlsx (V10.3) + một tệp Word tự tạo, đo thời gian nạp, rồi hỏi:
- câu có đáp án trong Excel (cơ quan thực hiện, thời hạn) -> câu trả lời chứa đáp án + nguồn đúng thủ tục
- câu có đáp án trong Word
- câu KHÔNG có trong dữ liệu -> phải nói không có, không nêu nguồn
Câu hỏi sinh tự động từ dữ liệu; đáp án = giá trị trong ô. Kết quả: system4/eval/results/nv3_specialist.json
"""
import json
import re
import sys
import tempfile
import time
from pathlib import Path

import httpx
import openpyxl

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8399"
XLSX = Path(r"D:/Claude/LLM for Procedures V10.3/Database/MAU_100_THU_TUC.xlsx")
OUT = Path(__file__).resolve().parent / "results" / "nv3_specialist.json"
c = httpx.Client(base_url=BASE, timeout=600)
user = f"spec{int(time.time())}"
assert c.post("/s4/auth/signup", json={"username": user, "password": "matkhau123"}).status_code == 200


def upload(path: Path) -> dict:
    t0 = time.time()
    with open(path, "rb") as f:
        r = c.post("/s4/datasets", files={"file": (path.name, f)})
    assert r.status_code == 200, r.text
    ds_id = r.json()["dataset"]["id"]
    while True:
        d = next(x for x in c.get("/s4/datasets").json()["datasets"] if x["id"] == ds_id)
        if d["status"] in ("ready", "error"):
            d["seconds"] = round(time.time() - t0, 1)
            return d
        time.sleep(1)


def chat(text):
    t0, first, evs = time.time(), None, []
    with c.stream("POST", "/s4/chat", json={"text": text}) as r:
        buf = ""
        for chunk in r.iter_text():
            buf += chunk
            while "\n\n" in buf:
                line, buf = buf.split("\n\n", 1)
                if line.startswith("data: "):
                    ev = json.loads(line[6:]); evs.append(ev)
                    if ev["type"] == "delta" and first is None:
                        first = time.time() - t0
    total = time.time() - t0
    msg = c.get(f"/s4/conversations/{evs[0]['conversation_id']}/messages").json()["messages"][-1]
    return msg, first or total, total


def norm(s):
    return re.sub(r"\s+", " ", (s or "").lower()).strip()


# ---- dữ liệu
docx_path = Path(tempfile.mkdtemp()) / "chinh_sach_team7.docx"
import docx
d = docx.Document()
d.add_heading("Chính sách đổi trả", level=1)
d.add_paragraph("Khách hàng được đổi trả sản phẩm trong vòng 7 ngày kể từ ngày nhận hàng nếu sản phẩm còn nguyên tem và hoá đơn.")
d.add_heading("Giờ làm việc", level=1)
d.add_paragraph("Bộ phận hỗ trợ làm việc từ 8 giờ đến 17 giờ 30, thứ Hai đến thứ Bảy. Chủ nhật nghỉ.")
t = d.add_table(rows=4, cols=2)
for i, (a, b) in enumerate([("Hạng thành viên", "Ưu đãi"), ("Bạc", "Giảm 5% mọi đơn hàng"), ("Vàng", "Giảm 10% mọi đơn hàng"), ("Kim cương", "Giảm 15% và miễn phí giao hàng")]):
    t.cell(i, 0).text, t.cell(i, 1).text = a, b
d.save(docx_path)

ing = [upload(XLSX), upload(docx_path)]
for x in ing:
    print(f"nạp {x['filename']}: {x['status']} · {x['n_records']} bản ghi · {x['seconds']}s · {x['message']}")

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
rows = list(wb["Thủ tục"].iter_rows(values_only=True))
h = rows[0]
procs = [dict(zip(h, r)) for r in rows[1:]]
procs = [p for p in procs if p.get("executing_agency") and p.get("processing_time_text")]
pick = procs[::max(1, len(procs) // 6)][:6]

cases = []
for p in pick:
    agency = p["executing_agency"].split(",")[0].split(" - ")[0].strip()
    cases.append(("cơ quan", f"Thủ tục \"{p['name']}\" do cơ quan nào thực hiện?", [agency], p["name"]))
    first_time = re.match(r"\s*([\d.,]+\s*ngày)", p["processing_time_text"])
    if first_time:
        cases.append(("thời hạn", f"Thời hạn giải quyết thủ tục \"{p['name']}\" là bao lâu?", [first_time.group(1)], p["name"]))
cases += [("Word", "Được đổi trả sản phẩm trong bao nhiêu ngày?", ["7 ngày"], "Chính sách đổi trả"),
          ("Word", "Thành viên hạng Vàng được ưu đãi gì?", ["10%"], "Vàng"),
          ("Word", "Bộ phận hỗ trợ làm việc mấy giờ?", ["8 giờ", "17 giờ 30"], "Giờ làm việc")]
NOT_IN = ["Giá vàng SJC hôm nay bao nhiêu?", "Ai là tác giả truyện Kiều?"]

results, times = [], []
for kind, q, keys, expect_title in cases:
    m, first, total = chat(q)
    times.append((first, total))
    ans = norm(m["content"])
    hit = all(norm(k) in ans for k in keys)
    titles = [s["title"] for s in m["meta"].get("sources", [])] + [s["title"] for s in m["meta"].get("consulted", [])]
    src_ok = any(norm(expect_title)[:60] in norm(t) or norm(t) in norm(expect_title) for t in titles)
    results.append({"kind": kind, "q": q, "expect": keys, "answer": m["content"][:500], "pass": hit, "source_ok": src_ok,
                    "sources": titles[:5], "first_s": round(first, 1), "total_s": round(total, 1)})
    print(f"{'PASS' if hit else 'FAIL'} {'nguồn✓' if src_ok else 'nguồn✗'} [{kind}] {first:.1f}s/{total:.1f}s {q[:70]}")
for q in NOT_IN:
    m, first, total = chat(q)
    times.append((first, total))
    ans = norm(m["content"])
    said_no = any(w in ans for w in ("không có", "không tìm thấy", "chưa có", "không chứa", "không đề cập"))
    no_src = not m["meta"].get("sources")
    results.append({"kind": "ngoài dữ liệu", "q": q, "answer": m["content"][:500], "pass": said_no and no_src, "source_ok": no_src,
                    "first_s": round(first, 1), "total_s": round(total, 1)})
    print(f"{'PASS' if said_no and no_src else 'FAIL'} [ngoài dữ liệu] {first:.1f}s/{total:.1f}s {q} -> {m['content'][:90]!r}")

ft = sorted(t[0] for t in times); tt = sorted(t[1] for t in times)
summary = {"answer_ok": sum(r["pass"] for r in results), "source_ok": sum(r["source_ok"] for r in results), "total": len(results),
           "first_median": ft[len(ft) // 2], "first_max": ft[-1], "total_median": tt[len(tt) // 2], "total_max": tt[-1],
           "ingest": [{k: x[k] for k in ("filename", "n_records", "seconds", "message")} for x in ing]}
print(f"\nĐúng đáp án: {summary['answer_ok']}/{len(results)} · nguồn đúng: {summary['source_ok']}/{len(results)} · "
      f"chữ đầu trung vị {summary['first_median']:.1f}s (tối đa {summary['first_max']:.1f}s) · cả câu trung vị {summary['total_median']:.1f}s")
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps({"summary": summary, "results": results}, ensure_ascii=False, indent=2), encoding="utf-8")

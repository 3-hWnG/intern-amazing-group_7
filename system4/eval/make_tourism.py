"""Dữ liệu du lịch GIẢ, ngắn, để thử pipeline Chuyên gia (đường nhanh + đường cũ) và tải lên web.

    python system4/eval/make_tourism.py          # ghi system4/eval/data/du_lich_gia.xlsx

Tên địa danh, tỉnh, đặc sản, mùa đẹp là kiến thức phổ thông. Khách sạn, tour, giá, số điện thoại, địa chỉ là BỊA (không phải doanh nghiệp thật).
Ba sheet: Điểm đến (12), Khách sạn (12), Tour (10). Ngày khởi hành là ô ngày thật của Excel, giá là số (định dạng #,##0),
cột "Hồ bơi" / "Còn chỗ" đánh dấu x — giống tệp người dùng hay có. Có chủ đích: hai khách sạn cùng tên "Biển Xanh" (hỏi lại),
hai tour Đà Lạt, một điểm đến vé "Miễn phí" (cột số lẫn chữ). Bộ câu mù: bench_cases.py (suite "tourism", split "blind"),
soạn TRƯỚC khi làm đường nhanh, không chỉnh theo kết quả.
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "data" / "du_lich_gia.xlsx"

PLACES = [  # tên, tỉnh/thành, vùng, mùa đẹp nhất, hoạt động nổi bật, đặc sản, giá vé tham quan
    ("Sa Pa", "Lào Cai", "Miền Bắc", "Tháng 9 – tháng 11", "Ngắm ruộng bậc thang, leo Fansipan", "Thắng cố, cá hồi", 150000),
    ("Vịnh Hạ Long", "Quảng Ninh", "Miền Bắc", "Tháng 4 – tháng 6", "Du thuyền ngắm đảo đá vôi, chèo kayak", "Chả mực", 290000),
    ("Ninh Bình (Tràng An)", "Ninh Bình", "Miền Bắc", "Tháng 5 – tháng 6", "Đi thuyền qua hang động, thăm chùa Bái Đính", "Dê núi, cơm cháy", 250000),
    ("Hội An", "Quảng Nam", "Miền Trung", "Tháng 2 – tháng 4", "Dạo phố cổ, thả đèn hoa đăng", "Cao lầu, bánh mì", 120000),
    ("Huế", "Thừa Thiên Huế", "Miền Trung", "Tháng 1 – tháng 4", "Tham quan Đại Nội, lăng tẩm, nghe ca Huế", "Bún bò Huế, bánh bèo", 200000),
    ("Đà Nẵng", "Đà Nẵng", "Miền Trung", "Tháng 3 – tháng 8", "Tắm biển Mỹ Khê, lên Bà Nà, xem cầu Rồng phun lửa", "Mì Quảng", "Miễn phí"),
    ("Phong Nha – Kẻ Bàng", "Quảng Bình", "Miền Trung", "Tháng 3 – tháng 8", "Khám phá động Phong Nha, động Thiên Đường", "Cháo canh", 150000),
    ("Đà Lạt", "Lâm Đồng", "Tây Nguyên", "Tháng 12 – tháng 3", "Ngắm đồi chè, vườn hoa, hồ Xuân Hương", "Bánh tráng nướng, dâu tây", 100000),
    ("Nha Trang", "Khánh Hòa", "Miền Trung", "Tháng 2 – tháng 8", "Lặn ngắm san hô, tắm bùn khoáng", "Bún chả cá", 80000),
    ("Phú Quốc", "Kiên Giang", "Miền Nam", "Tháng 11 – tháng 4", "Tắm biển Bãi Sao, câu mực đêm, cáp treo Hòn Thơm", "Gỏi cá trích, nước mắm", 50000),
    ("Cần Thơ", "Cần Thơ", "Miền Nam", "Tháng 9 – tháng 11", "Đi chợ nổi Cái Răng, vườn trái cây", "Lẩu mắm, bánh xèo", 30000),
    ("Côn Đảo", "Bà Rịa – Vũng Tàu", "Miền Nam", "Tháng 3 – tháng 9", "Lặn biển, xem rùa đẻ trứng, thăm nhà tù Côn Đảo", "Hạt bàng", 40000),
]
HOTELS = [  # tên, điểm đến, hạng sao, giá phòng/đêm, hồ bơi, số điện thoại (bịa), địa chỉ (bịa)
    ("Khách sạn Mây Ngàn", "Sa Pa", 4, 1450000, "", "0900 101 001", "12 Đường Fansipan, Sa Pa"),
    ("Khách sạn Sương Sớm", "Sa Pa", 3, 750000, "", "0900 101 002", "5 Đường Cầu Mây, Sa Pa"),
    ("Khách sạn Ngọc Vịnh", "Vịnh Hạ Long", 5, 2800000, "x", "0900 101 003", "88 Đường Hạ Long, Bãi Cháy"),
    ("Khách sạn Phố Hội", "Hội An", 4, 1200000, "x", "0900 101 004", "21 Trần Phú, Hội An"),
    ("Khách sạn Cố Đô", "Huế", 3, 650000, "", "0900 101 005", "9 Lê Lợi, Huế"),
    ("Khách sạn Sóng Bạc", "Đà Nẵng", 5, 3100000, "x", "0900 101 006", "150 Võ Nguyên Giáp, Đà Nẵng"),
    ("Khách sạn Đồi Thông", "Đà Lạt", 3, 700000, "", "0900 101 007", "3 Hồ Tùng Mậu, Đà Lạt"),
    ("Khách sạn Hoa Đào", "Đà Lạt", 4, 1100000, "x", "0900 101 008", "17 Trần Hưng Đạo, Đà Lạt"),
    ("Khách sạn Biển Xanh Nha Trang", "Nha Trang", 4, 1350000, "x", "0900 101 009", "45 Trần Phú, Nha Trang"),
    ("Khách sạn Biển Xanh Phú Quốc", "Phú Quốc", 5, 2950000, "x", "0900 101 010", "Bãi Trường, Phú Quốc"),
    ("Khách sạn Sông Hậu", "Cần Thơ", 3, 600000, "", "0900 101 011", "2 Hai Bà Trưng, Cần Thơ"),
    ("Khách sạn Hang Sáng", "Phong Nha – Kẻ Bàng", 2, 450000, "", "0900 101 012", "Thôn Sơn Trạch, Bố Trạch"),
]
TOURS = [  # mã, tên, điểm đến, số ngày, giá, ngày khởi hành, còn chỗ
    ("T01", "Sa Pa mùa lúa chín", "Sa Pa", 3, 3200000, date(2026, 9, 15), "x"),
    ("T02", "Hạ Long du thuyền 2 ngày", "Vịnh Hạ Long", 2, 2900000, date(2026, 10, 3), "x"),
    ("T03", "Tràng An – Bái Đính", "Ninh Bình (Tràng An)", 1, 990000, date(2026, 10, 25), ""),
    ("T04", "Di sản miền Trung", "Huế", 4, 5400000, date(2026, 11, 8), "x"),
    ("T05", "Hội An đêm hoa đăng", "Hội An", 2, 2300000, date(2026, 11, 20), "x"),
    ("T06", "Đà Lạt ngàn hoa", "Đà Lạt", 3, 2700000, date(2026, 12, 5), "x"),
    ("T07", "Đà Lạt săn mây", "Đà Lạt", 2, 1900000, date(2026, 12, 12), ""),
    ("T08", "Nha Trang biển đảo", "Nha Trang", 3, 3500000, date(2026, 12, 19), "x"),
    ("T09", "Phú Quốc nghỉ dưỡng", "Phú Quốc", 4, 6800000, date(2027, 1, 10), "x"),
    ("T10", "Miền Tây sông nước", "Cần Thơ", 2, 1600000, date(2027, 1, 24), ""),
]
SHEETS = [
    ("Điểm đến", ["Tên điểm đến", "Tỉnh/thành", "Vùng", "Mùa đẹp nhất", "Hoạt động nổi bật", "Đặc sản", "Giá vé tham quan (VNĐ)"], PLACES),
    ("Khách sạn", ["Tên khách sạn", "Điểm đến", "Hạng sao", "Giá phòng/đêm (VNĐ)", "Hồ bơi", "Số điện thoại", "Địa chỉ"], HOTELS),
    ("Tour", ["Mã tour", "Tên tour", "Điểm đến", "Số ngày", "Giá tour (VNĐ)", "Ngày khởi hành", "Còn chỗ"], TOURS),
]


def build(path: Path = OUT) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    head_font, body_font = Font(name="Arial", bold=True, color="FFFFFF"), Font(name="Arial")
    fill = PatternFill("solid", fgColor="1F4E78")
    for name, header, rows in SHEETS:
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, fill, Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
                if isinstance(c.value, int) and c.value >= 1000:
                    c.number_format = "#,##0"
                elif isinstance(c.value, date):
                    c.number_format = "dd/mm/yyyy"
        for i, h in enumerate(header, 1):
            width = max(len(str(h)), *(len(str(r[i - 1])) for r in rows))
            ws.column_dimensions[get_column_letter(i)].width = min(48, width + 3)
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    print(build())

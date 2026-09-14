"""Tệp thô -> danh sách chunk tìm kiếm được.

Nguyên tắc (slide 10 của bản brainstorm): phần còn lại của hệ thống KHÔNG cần
biết tệp gốc là csv, pdf hay docx. Ở đây quy mọi thứ về cùng một hình dạng:

    {"content": "...", "metadata": {...}}

Bảng biểu (csv/xlsx) chunk theo DÒNG — mỗi dòng là một đơn vị tra cứu độc lập,
đúng như "For a CSV, you might make each row a searchable unit". Văn bản dài
chunk theo đoạn, có phần chồng lấn để câu bị cắt ngang vẫn tìm được.
"""

from __future__ import annotations

import json
from pathlib import Path

from config import (ATTACH_CHUNK_CHARS, ATTACH_CHUNK_OVERLAP,
                    ATTACH_EXT_ALWAYS, ATTACH_EXT_OPTIONAL)
from domain.text import clean


class UnsupportedFile(Exception):
    pass


class MissingDependency(Exception):
    pass


def supported_extensions() -> set[str]:
    return set(ATTACH_EXT_ALWAYS) | set(ATTACH_EXT_OPTIONAL)


# --------------------------------------------------------------------------
# bảng biểu: mỗi DÒNG là một chunk
# --------------------------------------------------------------------------
def _rows_to_chunks(df, source: str) -> list[dict]:
    """Dòng -> "Cột: giá trị" xuống dòng. Giữ tên cột để câu hỏi khớp được."""
    chunks = []
    columns = [str(c) for c in df.columns]
    for idx, row in df.iterrows():
        pairs = []
        for col in columns:
            val = clean(row.get(col, ""))
            if val:
                pairs.append(f"{col}: {val}")
        if not pairs:
            continue
        chunks.append({
            "content": f"[Dòng {int(idx) + 1}]\n" + "\n".join(pairs),
            "metadata": {"kind": "row", "row": int(idx) + 1, "source": source},
        })
    return chunks


def _parse_table(path: Path, sep: str | None = None) -> list[dict]:
    import pandas as pd
    if path.suffix.lower() in {".xlsx", ".xls"}:
        df = pd.read_excel(path)
    else:
        df = pd.read_csv(path, sep=sep or ("\t" if path.suffix.lower() == ".tsv" else ","))
    df = df.fillna("")
    return _rows_to_chunks(df, path.name)


# --------------------------------------------------------------------------
# văn bản dài: chunk theo đoạn, có chồng lấn
# --------------------------------------------------------------------------
def split_text(text: str, source: str, *, page: int | None = None) -> list[dict]:
    text = (text or "").strip()
    if not text:
        return []

    size, overlap = ATTACH_CHUNK_CHARS, ATTACH_CHUNK_OVERLAP
    out, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            # lùi về ranh giới câu/đoạn gần nhất để không cắt giữa chừng
            window = text[start:end]
            for marker in ("\n\n", "\n", ". ", "; "):
                cut = window.rfind(marker)
                if cut > size * 0.5:
                    end = start + cut + len(marker)
                    break
        piece = text[start:end].strip()
        if piece:
            meta = {"kind": "text", "source": source}
            if page is not None:
                meta["page"] = page
            out.append({"content": piece, "metadata": meta})
        if end >= len(text):
            break                      # đã tới cuối: dừng, đừng phát lại phần đuôi
        next_start = end - overlap if overlap < size else end
        if next_start <= start:        # chống lặp vô hạn khi chunk quá ngắn
            next_start = end
        start = next_start
    return out


def _parse_txt(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return split_text(text, path.name)


def _parse_json(path: Path) -> list[dict]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        data = json.loads(raw)
    except Exception:
        return split_text(raw, path.name)

    if isinstance(data, list):
        out = []
        for i, item in enumerate(data):
            body = json.dumps(item, ensure_ascii=False, indent=2)
            out.append({"content": f"[Mục {i + 1}]\n{body}",
                        "metadata": {"kind": "row", "row": i + 1, "source": path.name}})
        return out
    return split_text(json.dumps(data, ensure_ascii=False, indent=2), path.name)


def _parse_pdf(path: Path) -> list[dict]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise MissingDependency(
            "Đọc PDF cần thư viện pypdf. Cài bằng:\n"
            "  .venv\\Scripts\\python.exe -m pip install pypdf"
        ) from exc
    reader = PdfReader(str(path))
    out = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        out.extend(split_text(text, path.name, page=i))
    if not out:
        raise UnsupportedFile(
            "PDF này không có lớp văn bản (nhiều khả năng là ảnh scan). "
            "Cần OCR trước khi tải lên.")
    return out


def _parse_docx(path: Path) -> list[dict]:
    try:
        import docx
    except ImportError as exc:
        raise MissingDependency(
            "Đọc DOCX cần thư viện python-docx. Cài bằng:\n"
            "  .venv\\Scripts\\python.exe -m pip install python-docx"
        ) from exc
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return split_text("\n".join(parts), path.name)


_HANDLERS = {
    ".csv": _parse_table,
    ".tsv": _parse_table,
    ".xlsx": _parse_table,
    ".xls": _parse_table,
    ".txt": _parse_txt,
    ".md": _parse_txt,
    ".json": _parse_json,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
}


def parse(path: Path) -> list[dict]:
    """Điểm vào duy nhất. Ném UnsupportedFile / MissingDependency khi không xử lý được."""
    ext = Path(path).suffix.lower()
    handler = _HANDLERS.get(ext)
    if handler is None:
        raise UnsupportedFile(
            f"Chưa hỗ trợ định dạng {ext or '(không rõ)'}. "
            f"Hỗ trợ: {', '.join(sorted(supported_extensions()))}")
    chunks = handler(Path(path))
    if not chunks:
        raise UnsupportedFile("Tệp rỗng hoặc không trích được nội dung.")
    return chunks


def describe(path: Path, chunks: list[dict]) -> str:
    """Một dòng mô tả cho manifest — mô hình đọc cái này để biết tệp chứa gì."""
    ext = Path(path).suffix.lower()
    kinds = {c["metadata"].get("kind") for c in chunks}
    if "row" in kinds:
        return f"bảng dữ liệu, {len(chunks)} dòng"
    pages = {c["metadata"].get("page") for c in chunks if c["metadata"].get("page")}
    if pages:
        return f"tài liệu {ext.lstrip('.')}, {len(pages)} trang, {len(chunks)} đoạn"
    return f"tài liệu {ext.lstrip('.')}, {len(chunks)} đoạn"

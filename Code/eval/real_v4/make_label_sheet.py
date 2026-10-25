"""Tệp Excel để gán nhãn bộ test real_v4 (cùng khuôn eval/real_v2/relabel/real_v2_relabel_B.xlsx). Lấy N câu ĐẦU theo thứ tự đã xáo
của cases_real_v4.json; gán từ trên xuống, dừng ở đâu thì bộ test là các câu đã gán liền mạch từ đầu.

    python eval/real_v4/make_label_sheet.py [N]      # mặc định 400 -> eval/real_v4/real_v4_label.xlsx
"""
import json
import os
import sqlite3
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(ROOT, "eval", "real_v2", "relabel"))
from make_relabel import BOLD, DB, FONT, GUIDE, INPUT_FILL, WRAP, fold, header  # noqa: E402

OUT = os.path.join(HERE, "real_v4_label.xlsx")
INTRO = [
    ("Gán nhãn bộ test real_v4", True),
    ("Mục đích: bộ test thứ hai cho vòng cải thiện System 3 trước khi nộp bài báo. Người sửa hệ thống (Claude) không đọc các câu này.", False),
    ("Gán TỪ TRÊN XUỐNG theo thứ tự, không bỏ câu. Dừng ở câu nào thì bộ test là các câu từ đầu tới đó (nên dừng ở bội số của 50).", False),
    ("Cách làm như lượt B của real_v2 (model pre-label rồi người kiểm từng câu) là được; ghi lại tên model và cách làm khi báo xong.", False),
    ("Hai điểm làm rõ (từ các câu bất đồng ở lượt gán lại real_v2, đúng theo REAL_V1_GUIDE.md mục 5):", True),
    ("  Kiến nghị, phản ánh, góp ý chính sách (kể cả không có dấu hỏi) -> apologize, không phải skip. skip chỉ cho câu cụt / không hiểu được.", False),
    ("  Hai thủ tục gần như tương đương, cái nào cũng đúng (ví dụ cấp bản sao trích lục / cấp lại) -> answer, ghi cả hai mã. clarify chỉ khi các thủ tục"
     " KHÁC NHAU về việc phải làm và câu hỏi không đủ để chọn.", False),
    ("", False),
]


def main(n: int):
    cases = json.load(open(os.path.join(HERE, "cases_real_v4.json"), encoding="utf-8"))[:n]
    wb = Workbook()
    g = wb.active
    g.title = "Hướng dẫn"
    g.column_dimensions["A"].width = 150
    lines = INTRO + [x for x in GUIDE[2:] if not x[0].startswith("Mục đích")]
    lines = [(t.replace("read_relabel.py rồi real_v1_agree.py để tính kappa và liệt kê câu bất đồng", "build_labeled.py để kiểm và khoá bộ test")
              .replace("eval/real_v2/real_v2_labels.json, eval/real_v2/cases_real_v2_labeled.json, eval/real_v2/_work/ (nhãn của người gán A)",
                       "kết quả hệ thống trên các câu này (chưa có; đừng chạy web Strict trên chúng)"), b) for t, b in lines]
    for r, (line, bold) in enumerate(lines, 1):
        c = g.cell(row=r, column=1, value=line)
        c.font, c.alignment = (BOLD if bold else FONT), Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("Nhãn")
    header(ws, [("id", 14), ("Câu hỏi (nguyên văn, đã che thông tin cá nhân)", 90), ("beh", 12), ("proc_id", 22), ("notes", 50)])
    dv = DataValidation(type="list", formula1='"answer,clarify,apologize,skip"', allow_blank=True)
    ws.add_data_validation(dv)
    for r, c in enumerate(cases, 2):
        ws.cell(row=r, column=1, value=c["id"]).font = FONT
        q = ws.cell(row=r, column=2, value=c["text"])
        q.font, q.alignment = FONT, WRAP
        for j in (3, 4, 5):
            x = ws.cell(row=r, column=j)
            x.fill, x.font, x.alignment = INPUT_FILL, FONT, WRAP
        dv.add(f"C{r}")
    cat = wb.create_sheet("Thủ tục")
    header(cat, [("proc_id", 14), ("name", 90), ("name_fold", 60), ("domain", 28), ("province", 16), ("levels", 18), ("in_scope", 9)])
    rows = sqlite3.connect(DB).execute("SELECT proc_id, name, domain, province, levels, in_scope FROM proc_flat").fetchall()
    rows.sort(key=lambda x: (-x[5], x[2] or "", x[1]))
    for r, (pid, name, dom, prov, lv, ins) in enumerate(rows, 2):
        for j, v in enumerate((pid, name, fold(name), dom, prov, lv, ins), 1):
            cat.cell(row=r, column=j, value=v).font = FONT
    cat.auto_filter.ref = f"A1:G{len(rows) + 1}"
    wb.active = 1
    wb.save(OUT)
    print(f"{len(cases)} câu -> {OUT}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(int(sys.argv[1]) if len(sys.argv) > 1 else 400)

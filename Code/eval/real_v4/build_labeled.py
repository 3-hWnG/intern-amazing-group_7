"""Đọc real_v4_label.xlsx đã gán -> real_v4_labels.json + cases_real_v4_labeled.json (bộ để chạy, bỏ skip), kiểm hợp lệ, in số đếm và sha256.

Bộ test = các câu đã gán LIỀN MẠCH từ câu đầu (thứ tự xáo của cases_real_v4.json); câu sau chỗ dừng bị bỏ qua. Không in câu hỏi.

    python eval/real_v4/build_labeled.py [tệp.xlsx]
"""
import hashlib
import json
import os
import sqlite3
import sys
from collections import Counter

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
DB = os.path.join(ROOT, "data", "runtime", "knowledge", "3-423eff02f8.db")
BEH = {"answer", "clarify", "apologize", "skip"}


def main(path: str):
    cases = json.load(open(os.path.join(HERE, "cases_real_v4.json"), encoding="utf-8"))
    by = {c["id"]: c for c in cases}
    known = {r[0] for r in sqlite3.connect(DB).execute("SELECT proc_id FROM proc_flat")}
    rows, errs = [], []
    for i, beh, pids, notes in ((r[0], r[2], r[3], r[4]) for r in load_workbook(path, read_only=True)["Nhãn"].iter_rows(min_row=2, values_only=True)):
        if not i:
            continue
        if not beh:
            break                                  # dừng ở câu chưa gán đầu tiên
        beh = str(beh).strip().lower()
        exp = [p.strip() for p in str(pids or "").replace(";", ",").split(",") if p.strip()]
        if beh not in BEH:
            errs.append(f"{i}: beh sai ({beh!r})")
        if beh == "answer" and not exp:
            errs.append(f"{i}: answer mà không có proc_id")
        if beh in ("apologize", "skip") and exp:
            errs.append(f"{i}: {beh} mà có proc_id")
        errs += [f"{i}: proc_id {p} không có trong DB bản 3" for p in exp if p not in known]
        rows.append({"id": i, "beh": beh, "expect_ids": exp, "fields": [], "notes": notes or ""})
    if [r["id"] for r in rows] != [c["id"] for c in cases[:len(rows)]]:
        errs.append("thứ tự id trong tệp khác cases_real_v4.json")
    if errs:
        sys.exit("Chưa hợp lệ:\n  " + "\n  ".join(errs))
    json.dump(rows, open(os.path.join(HERE, "real_v4_labels.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    run = [{"id": r["id"], "q": by[r["id"]]["text"], "beh": r["beh"], "expect_ids": r["expect_ids"], "fields": [], "notes": r["notes"],
            "stratum": "random", "source": by[r["id"]]["source"]} for r in rows if r["beh"] != "skip"]
    json.dump(run, open(os.path.join(HERE, "cases_real_v4_labeled.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"đã gán {len(rows)} câu liền mạch từ đầu; nhãn {dict(Counter(r['beh'] for r in rows))}; bộ chạy {len(run)} câu")
    for f in ("real_v4_labels.json", "cases_real_v4_labeled.json"):
        print(f, hashlib.sha256(open(os.path.join(HERE, f), "rb").read()).hexdigest())


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "real_v4_label.xlsx"))

"""Đọc real_v2_relabel_B.xlsx (người gán B đã điền) -> real_v2_labels_B.json, kiểm hợp lệ, rồi so với nhãn A (real_v2_labels.json) trên đúng 100 câu.

In Cohen's kappa trên beh (ba lớp, câu skip tính riêng) kèm khoảng tin cậy 95% bootstrap, % khớp expect_ids, danh sách câu bất đồng
(dùng lại eval/real_v1/real_v1_agree.py). Báo cả khi thấp. Nhãn A KHÔNG được sửa theo kết quả này; bất đồng chỉ ghi để báo cáo.

    python eval/real_v2/relabel/read_relabel.py [tệp.xlsx]
"""
import json
import os
import random
import sqlite3
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
V2 = os.path.dirname(HERE)
ROOT = os.path.dirname(os.path.dirname(V2))
sys.path.insert(0, os.path.join(ROOT, "eval", "real_v1"))
from real_v1_agree import BEH, kappa, report, s3  # noqa: E402

DB = os.path.join(ROOT, "data", "runtime", "knowledge", "3-423eff02f8.db")
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "real_v2_relabel_B.xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(XLSX)), "real_v2_labels_B.json")


def main():
    ids = json.load(open(os.path.join(HERE, "relabel_ids.json"), encoding="utf-8"))["ids"]
    known = {r[0] for r in sqlite3.connect(DB).execute("SELECT proc_id FROM proc_flat")}
    ws = load_workbook(XLSX, read_only=True)["Nhãn"]
    rows, errs = [], []
    for i, beh, pids, notes in ((r[0], r[2], r[3], r[4]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]):
        beh = (beh or "").strip().lower()
        expect = [p.strip() for p in str(pids or "").replace(";", ",").split(",") if p.strip()]
        if beh not in BEH + ("skip",):
            errs.append(f"{i}: beh trống hoặc sai ({beh!r})")
        if beh == "answer" and not expect:
            errs.append(f"{i}: answer mà không có proc_id")
        if beh in ("apologize", "skip") and expect:
            errs.append(f"{i}: {beh} mà có proc_id")
        errs += [f"{i}: proc_id {p} không có trong DB bản 3" for p in expect if p not in known]
        rows.append({"id": i, "beh": beh, "expect_ids": expect, "notes": notes or ""})
    if sorted(r["id"] for r in rows) != sorted(ids):
        errs.append("bộ id trong tệp khác relabel_ids.json")
    if errs:
        sys.exit("Chưa hợp lệ:\n  " + "\n  ".join(errs))
    json.dump(rows, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    A = {r["id"]: r for r in json.load(open(os.path.join(V2, "real_v2_labels.json"), encoding="utf-8")) if r["id"] in set(ids)}
    B = {r["id"]: r for r in rows}
    dis = s3(A, B)
    kept = [i for i in ids if "skip" not in (A[i]["beh"], B[i]["beh"])]
    rng = random.Random(20261025)
    boot = sorted(kappa(*zip(*[(A[j]["beh"], B[j]["beh"]) for j in (rng.choice(kept) for _ in kept)])) for _ in range(2000))
    print(f"kappa 95% bootstrap (2000 lần): {boot[50]:.3f} – {boot[1949]:.3f}")
    report(A, B, {r["id"]: r for r in json.load(open(os.path.join(V2, "cases_real_v2.json"), encoding="utf-8"))}, dis, False)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()

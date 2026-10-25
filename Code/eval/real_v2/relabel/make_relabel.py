"""Gói gán nhãn lại độc lập cho bộ test real_v2 (người gán thứ hai, B), để đo độ đồng thuận với nhãn gốc (A = Claude, README.md mục Nhãn).

Rút 100 câu ngẫu nhiên trong 313 câu của bộ chạy (cases_real_v2_labeled.json, chỉ đọc trường `id`), SEED cố định, rồi dựng
real_v2_relabel_B.xlsx: câu hỏi nguyên văn (từ cases_real_v2.json, tệp CHƯA có nhãn), cột để điền, danh mục 1.517 thủ tục bản 3 để tra.
Tệp KHÔNG chứa nhãn A. Chạy lại cùng SEED -> cùng 100 câu.

    python eval/real_v2/relabel/make_relabel.py
"""
import hashlib
import json
import os
import random
import sqlite3
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
V2 = os.path.dirname(HERE)
DB = os.path.join(V2, "..", "..", "data", "runtime", "knowledge", "3-423eff02f8.db")
SEED = 20261025
N = 100
OUT_IDS = os.path.join(HERE, "relabel_ids.json")
OUT_XLSX = os.path.join(HERE, "real_v2_relabel_B.xlsx")

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")

GUIDE = [
    ("Gán nhãn lại real_v2 (người gán B)", True),
    ("Mục đích: đo độ đồng thuận giữa hai người gán nhãn trên 100 câu rút ngẫu nhiên của bộ test real_v2 (bài báo cần Cohen's kappa).", False),
    ("", False),
    ("Trước khi làm, KHÔNG mở:", True),
    ("  eval/real_v2/real_v2_labels.json, eval/real_v2/cases_real_v2_labeled.json, eval/real_v2/_work/ (nhãn của người gán A)", False),
    ("  eval/hard/results/ và mục 5 của docs/BAO_CAO_CAU_KHO_2026-10-25.md (kết quả theo từng câu)", False),
    ("  Không chạy hệ thống (web Strict) trên các câu này, không hỏi Claude về câu cụ thể.", False),
    ("", False),
    ("Cách điền (sheet Nhãn, chỉ điền các ô nền vàng):", True),
    ("  beh: chọn answer / clarify / apologize / skip theo docs/eval 3 ngày/REAL_V1_GUIDE.md mục 5 (tóm tắt bên dưới).", False),
    ("  proc_id: bắt buộc khi answer, nhiều mã cách nhau bằng dấu phẩy. Khi clarify ghi các mã hợp lý nếu có. Để trống khi apologize / skip.", False),
    ("  notes: một dòng lý do. Không chép tên thật, SĐT, CCCD.", False),
    ("", False),
    ("Tra mã thủ tục:", True),
    ("  Sheet Thủ tục: lọc cột name_fold (chữ không dấu) bằng Text Filter > Contains. in_scope = 1 là 542 thủ tục cấp xã dùng chung toàn quốc (trả lời được).", False),
    ("  Hoặc: python eval/real_v2/find.py \"tam tru\" \"xac nhan\"  (thêm -a để xem cả bản theo tỉnh / ngoài cấp xã). Thử ít nhất ba cách viết trước khi kết luận không có.", False),
    ("", False),
    ("Tóm tắt quy tắc (REAL_V1_GUIDE.md mục 5):", True),
    ("  answer: chỉ một thủ tục in_scope = 1 (hoặc một cặp tương đương) rõ là cái người hỏi cần; nêu tỉnh có bản riêng thì lấy mã bản tỉnh đó.", False),
    ("  clarify: từ hai thủ tục khác nhau trở lên đều hợp lý mà câu không đủ để chọn; thủ tục chỉ có bản theo tỉnh mà chưa nêu tỉnh; hỏi thủ tục mà không nêu thủ tục nào.", False),
    ("  apologize: thủ tục không có trong dữ liệu, hoặc ngoài cấp xã; hỏi luật / xử phạt / tình trạng hồ sơ cá nhân / lỗi cổng; không phải câu hỏi thủ tục.", False),
    ("  skip: không phải câu hỏi, hoặc không hiểu được nếu không có ngữ cảnh gốc. Đừng dùng cho câu chỉ là khó.", False),
    ("  Hỏi tình trạng hồ sơ cụ thể -> apologize; hỏi thời gian giải quyết chung -> answer; hỏi luật mà dữ liệu thủ tục không ghi -> apologize.", False),
    ("", False),
    ("Ví dụ định dạng (không thuộc bộ):", True),
    ("  câu: \"Đăng ký tạm trú cần giấy tờ gì?\"  beh: answer  proc_id: <mã thủ tục đăng ký tạm trú tra được>  notes: dang ky tam tru, mot thu tuc ro", False),
    ("", False),
    ("Xong: lưu tệp này tại chỗ và báo Claude. Claude chạy read_relabel.py rồi real_v1_agree.py để tính kappa và liệt kê câu bất đồng.", False),
]


def fold(s: str) -> str:
    import unicodedata
    s = unicodedata.normalize("NFD", (s or "").lower()).replace("đ", "d")
    return "".join(ch for ch in s if unicodedata.category(ch) != "Mn")


def header(ws, cols):
    for j, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill, c.alignment = BOLD, HEAD_FILL, WRAP
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "A2"


def main():
    run_ids = sorted(x["id"] for x in json.load(open(os.path.join(V2, "cases_real_v2_labeled.json"), encoding="utf-8")))
    text = {c["id"]: c for c in json.load(open(os.path.join(V2, "cases_real_v2.json"), encoding="utf-8"))}
    ids = random.Random(SEED).sample(run_ids, N)
    json.dump({"seed": SEED, "n": N, "population": "cases_real_v2_labeled.json (313 câu)", "ids": ids,
               "strata": dict(Counter(text[i]["stratum"] for i in ids))},
              open(OUT_IDS, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    wb = Workbook()
    g = wb.active
    g.title = "Hướng dẫn"
    g.column_dimensions["A"].width = 150
    for r, (line, bold) in enumerate(GUIDE, 1):
        c = g.cell(row=r, column=1, value=line)
        c.font, c.alignment = (BOLD if bold else FONT), Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("Nhãn")
    header(ws, [("id", 14), ("Câu hỏi (nguyên văn, đã che thông tin cá nhân)", 90), ("beh", 12), ("proc_id", 22), ("notes", 50)])
    dv = DataValidation(type="list", formula1='"answer,clarify,apologize,skip"', allow_blank=True)
    ws.add_data_validation(dv)
    for r, i in enumerate(ids, 2):
        ws.cell(row=r, column=1, value=i).font = FONT
        q = ws.cell(row=r, column=2, value=text[i]["text"])
        q.font, q.alignment = FONT, WRAP
        for j in (3, 4, 5):
            c = ws.cell(row=r, column=j)
            c.fill, c.font, c.alignment = INPUT_FILL, FONT, WRAP
        dv.add(f"C{r}")

    cat = wb.create_sheet("Thủ tục")
    cols = [("proc_id", 14), ("name", 90), ("name_fold", 60), ("domain", 28), ("province", 16), ("levels", 18), ("in_scope", 9)]
    header(cat, cols)
    rows = sqlite3.connect(DB).execute("SELECT proc_id, name, domain, province, levels, in_scope FROM proc_flat").fetchall()
    rows.sort(key=lambda x: (-x[5], x[2] or "", x[1]))
    for r, (pid, name, dom, prov, lv, ins) in enumerate(rows, 2):
        for j, v in enumerate((pid, name, fold(name), dom, prov, lv, ins), 1):
            cat.cell(row=r, column=j, value=v).font = FONT
    cat.auto_filter.ref = f"A1:G{len(rows) + 1}"

    wb.active = 1
    wb.save(OUT_XLSX)
    print(f"{N} câu, strata {dict(Counter(text[i]['stratum'] for i in ids))}, danh mục {len(rows)} thủ tục")
    print("relabel_ids.json sha256", hashlib.sha256(open(OUT_IDS, "rb").read()).hexdigest())


if __name__ == "__main__":
    main()
